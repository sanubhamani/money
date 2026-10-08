"""Builds template/Budget-Template.xlsx. Run: python3 tools/make_template.py"""
import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

INK = "15151A"; SOFT = "F6F4EF"; VIOLET = "6A4CF0"; LINE = "E8E4DA"
head_font = Font(bold=True, color="FFFFFF", size=12)
head_fill = PatternFill("solid", fgColor=INK)
label_font = Font(bold=True, size=12, color=INK)
set_fill = PatternFill("solid", fgColor="ECE8FF")
note_font = Font(italic=True, size=10, color="6E6D78")
thin = Border(bottom=Side(style="thin", color=LINE))
MONEY = '#,##0'

def header(ws, row, names):
    for i, n in enumerate(names, 1):
        c = ws.cell(row=row, column=i, value=n)
        c.font = head_font; c.fill = head_fill
        c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[row].height = 26

wb = Workbook()
ws = wb.active
ws.title = "Budget"
ws["A1"] = "Snapshot Date"; ws["B1"] = datetime.date(2026, 10, 9)
ws["B1"].number_format = "yyyy-mm-dd"
ws["A2"] = "Currency"; ws["B2"] = "PKR"
for r in (1, 2):
    ws.cell(row=r, column=1).font = label_font
    ws.cell(row=r, column=2).fill = set_fill
    ws.cell(row=r, column=2).font = Font(bold=True, size=12, color=VIOLET)
ws["C1"] = "Change this each week to the day you update."
ws["C2"] = "3-letter code, e.g. PKR, USD, AED, GBP."
ws["C1"].font = note_font; ws["C2"].font = note_font
dv = DataValidation(type="list", formula1='"PKR,USD,AED,SAR,GBP,EUR,CAD,INR"', allow_blank=False,
                    showErrorMessage=False)  # list is a shortcut; other codes still allowed
ws.add_data_validation(dv); dv.add("B2")

header(ws, 4, ["Category", "Monthly Budget", "Spent So Far"])
rows = [("Food", 15000, 0), ("Transport", 6000, 0), ("Shopping", 8000, 0),
        ("Entertainment", 5000, 0), ("Bills", 12000, 0)]
for i, (n, b, s) in enumerate(rows, 5):
    ws.cell(row=i, column=1, value=n)
    ws.cell(row=i, column=2, value=b).number_format = MONEY
    ws.cell(row=i, column=3, value=s).number_format = MONEY
    for c in range(1, 4): ws.cell(row=i, column=c).border = thin
ws.cell(row=5, column=5, value="Add or rename categories below the header. Spent So Far = the month's total so far, not just this week.").font = note_font
ws.column_dimensions["A"].width = 22; ws.column_dimensions["B"].width = 18; ws.column_dimensions["C"].width = 18
ws.freeze_panes = "A5"

sv = wb.create_sheet("Savings")
header(sv, 1, ["Goal Name", "Target Amount", "Amount Saved", "Goal Deadline"])
goals = [("New laptop", 150000, 0, datetime.date(2027, 3, 31)), ("Emergency fund", 50000, 0, None)]
for i, (n, t, s, d) in enumerate(goals, 2):
    sv.cell(row=i, column=1, value=n)
    sv.cell(row=i, column=2, value=t).number_format = MONEY
    sv.cell(row=i, column=3, value=s).number_format = MONEY
    c = sv.cell(row=i, column=4, value=d); c.number_format = "yyyy-mm-dd"
    for col in range(1, 5): sv.cell(row=i, column=col).border = thin
sv.cell(row=2, column=6, value="Goal Deadline is optional. Amount Saved = the total in that goal today.").font = note_font
for col, w in zip("ABCD", (22, 18, 18, 18)): sv.column_dimensions[col].width = w
sv.freeze_panes = "A2"

wb.save("template/Budget-Template.xlsx")
print("ok")
